
from pathlib import Path

from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader,
    Docx2txtLoader,
    CSVLoader,
)
from langchain_core.documents import Document
from pptx import Presentation
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = [
    ".pdf",
    ".txt",
    ".docx",
    ".csv",
    ".md",
    ".pptx",
    ".xlsx",
]


def load_document(file_path):
    """Load supported files as LangChain Document objects."""

    path = Path(file_path)
    extension = path.suffix.lower()

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    if extension == ".pdf":
        return PyPDFLoader(str(path)).load()

    if extension == ".txt":
        return TextLoader(
            str(path), encoding="utf-8"
        ).load()

    if extension == ".docx":
        return Docx2txtLoader(str(path)).load()

    if extension == ".csv":
        return CSVLoader(
            str(path), encoding="utf-8"
        ).load()

    if extension == ".md":
        return TextLoader(
            str(path), encoding="utf-8"
        ).load()

    if extension == ".pptx":
        presentation = Presentation(str(path))
        documents = []

        for slide_number, slide in enumerate(
            presentation.slides, start=1
        ):
            slide_text = []

            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        slide_text.append(text)

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        values = [
                            cell.text.strip()
                            for cell in row.cells
                        ]
                        slide_text.append(" | ".join(values))

            if slide_text:
                documents.append(
                    Document(
                        page_content="\n".join(slide_text),
                        metadata={
                            "source": str(path),
                            "slide": slide_number,
                        },
                    )
                )

        if not documents:
            raise ValueError(
                "No readable text found in the PowerPoint file."
            )

        return documents

    if extension == ".xlsx":
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )
        documents = []

        try:
            for sheet in workbook.worksheets:
                rows = []

                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value).strip()
                        for value in row
                        if value is not None
                    ]
                    if values:
                        rows.append(" | ".join(values))

                if rows:
                    documents.append(
                        Document(
                            page_content="\n".join(rows),
                            metadata={
                                "source": str(path),
                                "sheet": sheet.title,
                            },
                        )
                    )
        finally:
            workbook.close()

        if not documents:
            raise ValueError(
                "No readable data found in the Excel file."
            )

        return documents

    raise ValueError(
        f"Unsupported file type: {extension}. "
        f"Supported types: {', '.join(SUPPORTED_EXTENSIONS)}"
    )
